import os
import logfire


def _parse_docx(file_path: str) -> str:
    import docx

    doc = docx.Document(file_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

    # Also capture table contents
    table_texts = []
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                table_texts.append(row_text)

    all_parts = paragraphs + table_texts
    return "\n".join(all_parts)


def _parse_pptx(file_path: str) -> str:
    from pptx import Presentation

    prs = Presentation(file_path)
    text_parts = []

    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = paragraph.text.strip()
                    if line:
                        text_parts.append(line)
            elif shape.has_table:
                for row in shape.table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        text_parts.append(row_text)

    return "\n".join(text_parts)


def parse_office(file_path: str) -> str:
    """
    Parses Office documents (.docx, .pptx) efficiently.
    Uses native python-docx and python-pptx, falling back to unstructured if needed.
    """
    with logfire.span("📄 Office Document Parsing", filename=file_path):
        ext = os.path.splitext(file_path)[1].lower()
        full_text = ""

        try:
            if ext == ".docx":
                full_text = _parse_docx(file_path)
            elif ext == ".pptx":
                full_text = _parse_pptx(file_path)

            if not full_text.strip():
                # Fallback to unstructured if native parser yielded no text
                try:
                    from unstructured.partition.auto import partition

                    elements = partition(filename=file_path)
                    full_text = "\n".join([str(el) for el in elements])
                except Exception as unstruct_err:
                    logfire.warning(f"Unstructured fallback failed for {file_path}: {unstruct_err}")

            if not full_text.strip():
                logfire.warning(f"⚠️ Empty text parsed for {file_path}")
            else:
                logfire.info(f"✅ Successfully parsed {len(full_text)} characters from {file_path}")

            return full_text

        except Exception as e:
            logfire.error(f"❌ Office Parse Failed for {file_path}: {e}")
            raise e
